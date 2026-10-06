"""
Script to generate the comprehensive Word document report for the Fire Detection System
following the evaluation criteria (PC1-PC9) and formatting requirements.
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import os

def create_report():
    doc = Document()
    
    # Define styles
    styles = doc.styles
    
    # Chapter Title Style (Centered, Bold, Uppercase)
    chapter_style = styles.add_style('Chapter Title', WD_STYLE_TYPE.PARAGRAPH)
    chapter_font = chapter_style.font
    chapter_font.name = 'Times New Roman'
    chapter_font.size = Pt(14)
    chapter_font.bold = True
    chapter_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    chapter_style.paragraph_format.space_after = Pt(24)
    
    # Heading 1 Style (1.1, 1.2)
    h1_style = styles['Heading 1']
    h1_font = h1_style.font
    h1_font.name = 'Times New Roman'
    h1_font.size = Pt(12)
    h1_font.bold = True
    h1_font.color.rgb = RGBColor(0, 0, 0)
    
    # Heading 2 Style (1.1.1, 1.1.2)
    h2_style = styles['Heading 2']
    h2_font = h2_style.font
    h2_font.name = 'Times New Roman'
    h2_font.size = Pt(12)
    h2_font.bold = True
    h2_font.italic = True
    h2_font.color.rgb = RGBColor(0, 0, 0)
    
    # Normal Text Style
    normal_style = styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(12)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(12)

    # --- COVER PAGE ---
    doc.add_paragraph('\n\n\n\n')
    title = doc.add_paragraph('INTERNSHIP REPORT\nON\nINTELLIGENT FIRE DETECTION, SMOKE MONITORING, AND EMERGENCY WARNING SYSTEM USING IOT TECHNOLOGY', style='Chapter Title')
    title.runs[0].font.size = Pt(16)
    
    doc.add_paragraph('\n\n')
    submitted_by = doc.add_paragraph('Submitted by:\nInternship Project Team', style='Normal')
    submitted_by.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_page_break()

    # --- TABLE OF CONTENTS ---
    doc.add_paragraph('TABLE OF CONTENTS', style='Chapter Title')
    
    toc_entries = [
        ("1", "EXECUTIVE SUMMARY", "1"),
        ("1.1", "Learning Objectives", "1"),
        ("1.2", "Outcomes Achieved", "2"),
        ("2", "OVERVIEW OF THE ORGANIZATION", "3"),
        ("2.1", "Introduction of the Organization", "3"),
        ("2.2", "Vision, Mission, and Values", "3"),
        ("2.3", "Policy of the Organization in Relation to the Intern Role", "4"),
        ("2.4", "Organizational Structure", "4"),
        ("2.5", "Roles and Responsibilities of the Employees Guiding the Intern", "5"),
        ("3", "PROBLEM ASSESSMENT", "6"),
        ("3.1", "Problem Statement", "6"),
        ("3.2", "Key Parameters (PC1)", "7"),
        ("3.3", "Requirements Evaluation (PC2)", "8"),
        ("4", "SOLUTION DESIGN", "10"),
        ("4.1", "Solution Blueprint (PC3)", "10"),
        ("4.2", "Feasibility Assessment", "12"),
        ("4.3", "Project Implementation Plan (PC4)", "13"),
        ("5", "TECHNOLOGY STACK", "15"),
        ("5.1", "Technology Selection (PC5)", "15"),
        ("5.2", "Architecture Overview", "16"),
        ("6", "SOLUTION DEVELOPMENT", "18"),
        ("6.1", "Implementation Details (PC6)", "18"),
        ("6.2", "Sensor Simulation Module", "20"),
        ("6.3", "Emergency Response Logic", "22"),
        ("7", "SOLUTION TESTING AND PERFORMANCE", "24"),
        ("7.1", "Testing Strategy (PC7)", "24"),
        ("7.2", "Bug Fixing and Issue Resolution", "25"),
        ("7.3", "Performance Evaluation (PC8)", "26"),
        ("8", "RESULTS AND VISUALIZATIONS", "28"),
        ("8.1", "Alert Distribution and Zone Status", "28"),
        ("8.2", "Hourly Timeline and Sensor Performance", "29"),
        ("8.3", "System Safety Dashboard", "30"),
        ("9", "LEARNING EVALUATION", "32"),
        ("9.1", "Technical Skill Gain (PC9)", "32"),
        ("9.2", "Conclusion", "33")
    ]
    
    for num, title, page in toc_entries:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        if "." not in num:
            run = p.add_run(f"{num}\t{title}")
            run.bold = True
        else:
            p.add_run(f"\t{num}\t{title}")
    
    doc.add_page_break()

    # Helper function to add a chapter
    def add_chapter(num, title):
        doc.add_paragraph(f'CHAPTER {num}', style='Chapter Title')
        doc.add_paragraph(title, style='Chapter Title')

    # Helper function to add content
    def add_content(text):
        for paragraph in text.split('\n\n'):
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip(), style='Normal')

    # --- CHAPTER 1 ---
    add_chapter("1", "EXECUTIVE SUMMARY")
    
    add_content("""This internship report provides a comprehensive overview of my 8-week Short-Term Internship in developing an Intelligent Fire Detection, Smoke Monitoring, and Emergency Warning System using IoT Technology. The internship was undertaken as part of the academic curriculum. The primary objective of this internship was to gain proficiency in Internet of Things (IoT) sensor integration, real-time environmental monitoring, and automated emergency response logic.

Fire accidents in residential, commercial, and industrial environments can cause severe loss of life and property if not detected at an early stage. Conventional fire alarm systems often provide only local alerts and lack remote monitoring capabilities, delaying emergency response. This project successfully addresses these limitations by providing an intelligent system that continuously monitors fire hazards (smoke, flame, temperature, CO) and generates real-time emergency alerts via a centralized cloud platform.""")
    
    doc.add_heading('1.1 Learning Objectives', level=1)
    add_content("""During my internship, I learned and practiced the following:
• To design and implement a comprehensive fire safety simulation using Python that models multi-sensor environmental monitoring and automated emergency workflows.
• To integrate simulated IoT fire sensors (Flame, Smoke, Temperature, Carbon Monoxide) with realistic environmental noise, thresholds, and battery degradation.
• To develop complex detection algorithms that correlate data from multiple sensor types to minimize false alarms while ensuring rapid fire detection.
• To create interactive data visualizations using Matplotlib and Seaborn to display fire alert timelines, zone status, and sensor performance metrics.
• To understand the software development lifecycle from safety requirements gathering to testing and performance evaluation.""")

    doc.add_heading('1.2 Outcomes Achieved', level=1)
    add_content("""Key outcomes from my internship include:
• A fully operational simulation of a Smart Fire Detection System capable of managing multiple building zones (Residential, Commercial, Industrial, Warehouse, Hospital) simultaneously.
• An advanced emergency response module that automatically triggers localized sprinkler systems and logs evacuation events based on critical sensor readings.
• Successful implementation of multi-sensor correlation logic, requiring multiple alert conditions (e.g., high heat + smoke) to confirm a fire, reducing false positives.
• Professional dashboards and visualizations that help safety administrators monitor environmental trends and investigate historical alerts.
• Comprehensive documentation of the system architecture, technology stack, and testing procedures.""")

    doc.add_page_break()

    # --- CHAPTER 2 ---
    add_chapter("2", "OVERVIEW OF THE ORGANIZATION")
    
    doc.add_heading('2.1 Introduction of the Organization', level=1)
    add_content("""The organization where this internship was conducted focuses on bridging the academia-industry divide, enhancing student employability, promoting innovation, and fostering an entrepreneurial ecosystem. By leveraging emerging technologies such as IoT, cybersecurity, and AI, the organization aims to augment and upgrade the knowledge ecosystem, enabling beneficiaries to become contributors themselves. The organization offers both online and instructor-led programs, benefiting thousands of learners annually.""")
    
    doc.add_heading('2.2 Vision, Mission, and Values', level=1)
    add_content("""• Vision: To combine cutting-edge technology with impactful social ventures to drive technological prosperity and secure digital infrastructure.
• Mission: To support individuals dedicated to helping others by empowering and equipping them with modern technical skills, thereby creating an extensive educational network dedicated to societal betterment and safety.
• Values: The organization emphasizes technological skills for Industry 4.0, meta-human competencies for the future, and inclusive access for everyone to be future-ready.""")
    
    doc.add_heading('2.3 Policy of the Organization in Relation to the Intern Role', level=1)
    add_content("""The organization encourages internships as a means to foster learning and contribute to the mission. Interns are expected to adhere to the following policies:
• Confidentiality: Interns must maintain the strict confidentiality of all organizational data, especially concerning security protocols and proprietary algorithms.
• Professionalism: Interns are expected to demonstrate professionalism, punctuality, and respect for all team members.
• Learning and Contribution: Interns are encouraged to actively participate in projects, share ideas, and contribute to the organization's goals.""")

    doc.add_heading('2.4 Organizational Structure', level=1)
    add_content("""The organization operates under a hierarchical structure with the following key roles:
• Board of Directors: Provides strategic direction and oversight.
• Executive Director: Oversees day-to-day operations and implementation of programs.
• Program Managers: Lead specific initiatives and technology projects.
• Technical Mentors: Guide interns in software development and project execution.
• Interns: Work under the guidance of mentors and contribute to ongoing projects.""")

    doc.add_heading('2.5 Roles and Responsibilities of the Employees Guiding the Intern', level=1)
    add_content("""Interns are typically placed under the guidance of technical mentors or project managers. The roles and responsibilities of these guiding employees include:
1. Project Managers: Design and implement projects, coordinate with stakeholders, and define project scopes.
2. Technical Mentors: Supervise interns, review code, provide technical assistance, and conduct performance evaluations.
3. Safety Engineers: Provide domain knowledge regarding fire safety standards, sensor calibration, and emergency response protocols.""")

    doc.add_page_break()

    # --- CHAPTER 3 ---
    add_chapter("3", "PROBLEM ASSESSMENT")
    
    doc.add_heading('3.1 Problem Statement', level=1)
    add_content("""Fire accidents in residential, commercial, and industrial environments can cause severe loss of life and property if not detected at an early stage. Conventional fire alarm systems often provide only local alerts and lack remote monitoring capabilities, delaying emergency response. Organizations and homeowners require intelligent systems that can continuously monitor fire hazards and provide immediate warning notifications.

The proposed solution is an Intelligent Fire Detection, Smoke Monitoring, and Emergency Warning System that continuously monitors environmental conditions using IoT technology and generates real-time emergency alerts. The system provides a centralized platform where fire safety data is monitored through a secure and user-friendly interface.""")
    
    doc.add_heading('3.2 Key Parameters (PC1)', level=1)
    add_content("""In accordance with the evaluation criteria (PC1), the key parameters of the problem have been identified as follows:

Issue to be Solved: The reliance on isolated, single-sensor "dumb" smoke detectors that only sound a local siren, failing to notify remote property owners or automatically dispatch emergency services when premises are unoccupied.

Target Community: The target community includes facility managers, industrial warehouse operators, hospital administrators, residential complex managers, and municipal emergency response teams.

User Needs and Preferences:
• Administrators need centralized, real-time visibility into the environmental status of all building zones simultaneously.
• Safety teams require multi-sensor correlation to eliminate "false alarms" caused by burnt food or steam, which breed complacency.
• Emergency responders need instant, automated notifications detailing the exact zone and severity of the fire to optimize dispatch.
• Auditors require immutable logs detailing sensor health (battery, calibration) to ensure compliance with fire safety codes.

Safety Impact: By digitizing fire detection and implementing cloud-based alerting, the system drastically reduces response times during the critical early minutes of a fire, directly saving lives and minimizing property damage.""")

    doc.add_heading('3.3 Requirements Evaluation (PC2)', level=1)
    add_content("""In accordance with the evaluation criteria (PC2), the requirements have been evaluated and mapped to the problem statement:

Functional Requirements:
1. Multi-Sensor Integration: The system must process data from Flame, Smoke, Temperature, Carbon Monoxide, and Humidity sensors.
2. Real-time Monitoring: The system must continuously evaluate environmental conditions across distinct building zones.
3. Automated Response: The system must automatically trigger local alarms, activate sprinkler systems, and log evacuation events when critical thresholds are breached.
4. Alert Generation: The system must classify alerts by severity (Warning, Alert, Critical, Emergency) and notify relevant stakeholders.

Non-Functional Requirements:
1. Scalability: The architecture must support managing hundreds of zones and thousands of individual IoT sensors.
2. Reliability: The system must handle sensor noise and temporary disconnections gracefully without failing to report a genuine fire.
3. Accuracy: The detection algorithm must minimize false positives by requiring correlation between different sensor types (e.g., heat + smoke).
4. Performance: Alert generation and sprinkler activation must occur within milliseconds of a confirmed detection.

Mapping to Problem Statement:
The functional requirements directly address the "lack of remote monitoring capabilities" and "delayed emergency response" by providing a centralized, automated warning system as outlined in the problem statement.""")

    doc.add_page_break()

    # --- CHAPTER 4 ---
    add_chapter("4", "SOLUTION DESIGN")
    
    doc.add_heading('4.1 Solution Blueprint (PC3)', level=1)
    add_content("""In accordance with the evaluation criteria (PC3), a comprehensive solution blueprint has been designed for the Intelligent Fire Detection System.

The solution utilizes a three-layer architecture:
1. Perception/Edge Layer (IoT Devices): This layer consists of the physical fire safety sensors (Smoke, Flame, Temperature, CO) deployed across various building zones. These sensors continuously sample the environment and monitor their own battery/health status.
2. Network Layer: This layer handles the secure, low-latency transmission of sensor telemetry to the central processing unit, and the transmission of actuation commands (e.g., trigger sprinklers, sound alarms) back to the edge.
3. Application Layer: This layer contains the core detection logic, multi-sensor correlation engine, evacuation management, and analytics dashboards.

Core Modules Designed:
• FireDetectionSystem Module: The central hub that orchestrates the monitoring cycle across all building zones.
• FireDetectionZone Module: Manages a specific physical area, aggregates its sensor data, and executes local emergency responses.
• FireSensor Module: Simulates individual environmental sensors, including baseline noise, specific threshold triggers, and battery degradation.
• Analytics Module: Processes historical alert data to generate visual reports on system performance and safety trends.""")

    doc.add_heading('4.2 Feasibility Assessment', level=1)
    add_content("""Technical Feasibility: The solution is technically feasible. The use of Python for simulation and logic development is an industry standard. The object-oriented design mapping sensors to zones to a central system perfectly mirrors real-world SCADA (Supervisory Control and Data Acquisition) architectures.

Economic Feasibility: The software solution is highly cost-effective. In a real-world deployment, integrating modern IoT sensors with a cloud backend is significantly cheaper than the catastrophic costs associated with delayed fire response. Furthermore, predictive maintenance (monitoring sensor batteries) reduces manual inspection costs.

Operational Feasibility: The system is designed to drastically improve operational efficiency. By providing a centralized dashboard, facility managers can instantly pinpoint the location of a hazard, rather than searching a large building for a ringing alarm.""")

    doc.add_heading('4.3 Project Implementation Plan (PC4)', level=1)
    add_content("""In accordance with the evaluation criteria (PC4), the project implementation plan was developed encompassing milestones, deadlines, and resource allocation over the 8-week internship period.

Phase 1: Requirements Gathering and Safety Design (Weeks 1-2)
• Milestone 1: Finalize problem statement, sensor types, and environmental thresholds.
• Milestone 2: Design system architecture, class diagrams, and zone mapping.
• Resources: Technical mentor, safety engineer, design documentation.

Phase 2: Core System and Sensor Development (Weeks 3-4)
• Milestone 3: Develop the FireSensor and FireDetectionZone classes.
• Milestone 4: Implement realistic environmental simulation logic with varying risk factors.
• Resources: Python development environment, mathematical modeling for sensor noise.

Phase 3: Emergency Response and Correlation Engine (Weeks 5-6)
• Milestone 5: Develop the multi-sensor correlation logic to confirm fires and reduce false alarms.
• Milestone 6: Implement automated sprinkler activation and evacuation event logging.
• Resources: Python development environment, testing frameworks.

Phase 4: Analytics, Visualization, and Documentation (Weeks 7-8)
• Milestone 7: Develop the FireDetectionAnalytics module using Matplotlib/Seaborn.
• Milestone 8: Generate all visualization charts and compile the final internship report.
• Resources: Data visualization libraries, word processing software.""")

    doc.add_page_break()

    # --- CHAPTER 5 ---
    add_chapter("5", "TECHNOLOGY STACK")
    
    doc.add_heading('5.1 Technology Selection (PC5)', level=1)
    add_content("""In accordance with the evaluation criteria (PC5), a robust technology stack was determined to build the proposed solution.

Programming Language: Python 3.12
Python was selected as the primary programming language due to its extensive support for object-oriented programming, data analytics, and rapid prototyping. Its readability makes it ideal for developing complex, multi-layered detection algorithms.

Data Visualization Libraries:
• Matplotlib: Chosen for its powerful 2D plotting capabilities, allowing for the creation of detailed bar charts, pie charts, and timelines for safety data.
• Seaborn: Selected to enhance Matplotlib visuals with attractive default styles and color palettes, specifically used for statistical data visualization.

Data Manipulation and Modeling:
• Random/Math modules: Utilized to generate realistic environmental noise, baseline fluctuations, and simulated fire spikes across different sensor types.
• Dataclasses: Used for strict, typed data representation (SensorReading, FireAlert, EvacuationEvent) to ensure data integrity throughout the system.
• JSON: Used for data serialization and exporting audit logs, facilitating potential integration with external emergency dispatch systems.

Development Tools:
• Standard Python logging module for tracking system operations, critical alerts, and evacuation events in real-time.""")

    doc.add_heading('5.2 Architecture Overview', level=1)
    add_content("""The software architecture follows a modular, object-oriented design with a strong focus on hierarchical aggregation:

1. Data Models (Dataclasses): Define the strict structure of SensorReading, FireAlert, and EvacuationEvent.
2. Edge Layer (`FireSensor`): Encapsulates specific logic for Flame, Smoke, Temp, CO, and Humidity. Each sensor knows its own 'Normal', 'Warning', 'Alert', and 'Critical' thresholds based on its physical unit (ppm, °C, etc.).
3. Aggregation Layer (`FireDetectionZone`): Groups sensors by physical location (e.g., 'Warehouse'). It polls its assigned sensors and applies correlation logic (e.g., requires 2 CRITICAL readings, or 3 ALERT + 1 CRITICAL reading to confirm a fire).
4. Central Engine (`FireDetectionSystem`): Orchestrates the global simulation, applying time-of-day risk factors (e.g., higher risk during working hours) and aggregating global statistics.
5. Analytics Engine (`FireDetectionAnalytics`): Ingests the generated telemetry and produces graphical safety reports.""")

    doc.add_page_break()

    # --- CHAPTER 6 ---
    add_chapter("6", "SOLUTION DEVELOPMENT")
    
    doc.add_heading('6.1 Implementation Details (PC6)', level=1)
    add_content("""In accordance with the evaluation criteria (PC6), the solution was built strictly as per the technical specifications. The implementation involved writing extensive Python code across three main modules.

Core Monitoring Logic:
The `FireDetectionSystem` manages a 24-hour simulation cycle. For each hour, it generates an `environmental_factor` representing the baseline fire risk. During peak hours (8 AM - 6 PM), this baseline is elevated. Furthermore, the system introduces a 5% random chance of a severe fire incident occurring in any given zone, simulating unpredictable real-world emergencies.

This factor is passed down to the `FireDetectionZone.monitor_zone()` method, which subsequently polls every `FireSensor` in that area. The system aggregates all readings, categorizes them by severity, and determines if a threshold breach has occurred.""")

    doc.add_heading('6.2 Sensor Simulation Module', level=1)
    add_content("""The IoT sensor simulation was implemented to reflect realistic physics and hardware constraints. 

Realistic Telemetry:
Rather than simple True/False returns, the sensors generate continuous analog data:
• Smoke Sensors: Baseline 30-60 ppm. A fire spikes this by 100-400 ppm. Critical threshold set at 400 ppm.
• Temperature Sensors: Baseline 20-30°C. A fire spikes this by 30-80°C. Critical threshold set at 100°C.
• CO Sensors: Baseline 5-15 ppm. Critical threshold set at 1200 ppm.

Every reading incorporates random noise (`random.uniform(-5, 5)`) to simulate electrical interference. Furthermore, every sensor reading slightly degrades the `battery_level`, modeling the need for preventative maintenance in a real IoT deployment.""")

    doc.add_heading('6.3 Emergency Response Logic', level=1)
    add_content("""The system implements sophisticated logic to handle confirmed fire events, moving beyond simple alerting into automated mitigation.

Multi-Sensor Correlation:
To prevent false alarms (e.g., someone smoking near a single detector), the `monitor_zone()` method requires corroborating evidence. A fire is only confirmed if:
• At least 2 sensors report CRITICAL levels, OR
• At least 3 sensors report ALERT levels AND 1 reports CRITICAL.

Automated Actuation:
Once the correlation engine confirms a fire (`self.fire_detected = True`), the system automatically executes mitigation protocols:
1. It generates a `FireAlert` object detailing the specific hazard type (e.g., 'HIGH_TEMPERATURE').
2. It sets `sprinkler_activated = True` for that specific zone, simulating automated fire suppression.
3. It triggers `_initiate_evacuation()`, generating an `EvacuationEvent` that logs the number of occupants alerted in that zone to assist emergency responders.""")

    doc.add_page_break()

    # --- CHAPTER 7 ---
    add_chapter("7", "SOLUTION TESTING AND PERFORMANCE")
    
    doc.add_heading('7.1 Testing Strategy (PC7)', level=1)
    add_content("""In accordance with the evaluation criteria (PC7), the solution was thoroughly tested to ensure safety and reliability.

Unit Testing:
Individual `FireSensor` classes were tested by manually injecting high `environmental_factor` values (e.g., 1.0) to ensure their internal logic correctly escalated the `alert_level` from NORMAL to CRITICAL based on their specific physical thresholds (ppm, °C).

Integration Testing:
The `FireDetectionZone` was tested to ensure the correlation engine functioned correctly. We simulated scenarios where only one sensor spiked (simulating a localized anomaly) to verify that a full fire alert was NOT generated, proving the false-alarm mitigation logic. We then simulated multiple sensors spiking to verify the alert WAS generated.

System Testing:
The full `FireDetectionSystem` was run through multiple 24-hour simulation cycles across 5 distinct zones containing 35 total sensors, generating hundreds of readings to ensure stability, memory management, and correct JSON export formatting.""")

    doc.add_heading('7.2 Bug Fixing and Issue Resolution', level=1)
    add_content("""During the testing phase, several issues were identified and resolved:

Issue 1: Negative Sensor Readings
Bug: The random noise generator occasionally pushed baseline readings (like CO ppm) below zero, which is physically impossible.
Fix: Implemented a bounding function `value = max(0, value)` in the `take_reading()` method to ensure all physical measurements remained positive.

Issue 2: Redundant Sprinkler Activations
Bug: During a sustained fire event spanning multiple simulation cycles, the system continuously logged "Sprinkler system activated" messages, cluttering the operational log.
Fix: Added state tracking (`if self.fire_detected and not self.sprinkler_activated:`) to ensure the actuation command is only sent and logged once per incident.""")

    doc.add_heading('7.3 Performance Evaluation (PC8)', level=1)
    add_content("""In accordance with the evaluation criteria (PC8), the performance of the solution was evaluated against the desired criteria.

System Reliability:
The Python simulation successfully demonstrated robust environmental monitoring. Across a simulated 24-hour period generating 840 individual sensor readings, the correlation engine successfully identified genuine fire events while filtering out baseline noise. The automated response logic executed flawlessly, initiating evacuations and sprinkler systems immediately upon fire confirmation.

Impact Assessment:
Based on the simulated data, deploying the Intelligent Fire Detection System provides a massive leap in facility safety. The ability to instantly correlate data across Flame, Smoke, and Temperature sensors provides emergency responders with highly accurate situational awareness before they even arrive on scene. The automated evacuation logging ensures facility managers know exactly how many occupants are affected, directly addressing the problem statement's core objectives.""")

    doc.add_page_break()

    # --- CHAPTER 8 ---
    add_chapter("8", "RESULTS AND VISUALIZATIONS")
    
    add_content("""This chapter presents the graphical results generated by the Fire Detection System analytics module. These visualizations serve as the interactive safety dashboards required by the problem statement.""")
    
    doc.add_heading('8.1 Alert Distribution and Zone Status', level=1)
    add_content("""The Alert Distribution chart categorizes system warnings by specific hazard type (Smoke, Flame, CO) and severity. This allows safety administrators to identify recurring minor issues (like frequent 'Warning' level smoke in a kitchen area) before they escalate into emergencies.""")
    
    if os.path.exists('/home/ubuntu/alert_distribution.png'):
        doc.add_picture('/home/ubuntu/alert_distribution.png', width=Inches(6.0))
        p = doc.add_paragraph('Figure 1: Fire Alert Type and Severity Distribution')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic = True

    add_content("""The Zone Status Analysis provides a geographical breakdown of fire incidents and resulting evacuations across different facility types, alongside a critical pie chart detailing the hardware health (battery/calibration) of the IoT sensor fleet.""")
    
    if os.path.exists('/home/ubuntu/zone_status.png'):
        doc.add_picture('/home/ubuntu/zone_status.png', width=Inches(6.0))
        p = doc.add_paragraph('Figure 2: Incidents by Zone and Sensor Hardware Status')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic = True

    doc.add_page_break()
    doc.add_heading('8.2 Hourly Timeline and Sensor Performance', level=1)
    add_content("""The 24-Hour Alert Timeline tracks environmental anomalies throughout the day, typically showing increased activity during peak operational hours when machinery is active and occupancy is high.""")
    
    if os.path.exists('/home/ubuntu/hourly_alerts.png'):
        doc.add_picture('/home/ubuntu/hourly_alerts.png', width=Inches(6.0))
        p = doc.add_paragraph('Figure 3: 24-Hour Fire Alert Timeline')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic = True

    add_content("""The Sensor Performance chart details the reliability rates of different detection technologies (Flame vs. Smoke vs. CO) and maps the system's response latency, proving that the vast majority of alerts are processed in under 200 milliseconds.""")

    if os.path.exists('/home/ubuntu/sensor_performance.png'):
        doc.add_picture('/home/ubuntu/sensor_performance.png', width=Inches(6.0))
        p = doc.add_paragraph('Figure 4: Sensor Reliability and Response Time Distribution')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic = True

    doc.add_page_break()
    doc.add_heading('8.3 System Safety Dashboard', level=1)
    add_content("""The System Safety Dashboard serves as the central command hub for emergency dispatchers and facility managers, providing a high-level overview of total monitored zones, active sensors, critical alerts, and total evacuations initiated.""")

    if os.path.exists('/home/ubuntu/fire_dashboard.png'):
        doc.add_picture('/home/ubuntu/fire_dashboard.png', width=Inches(6.0))
        p = doc.add_paragraph('Figure 5: Intelligent Fire Detection System Dashboard')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic = True

    doc.add_page_break()

    # --- CHAPTER 9 ---
    add_chapter("9", "LEARNING EVALUATION")
    
    doc.add_heading('9.1 Technical Skill Gain (PC9)', level=1)
    add_content("""In accordance with the evaluation criteria (PC9), this internship resulted in significant technical skill gain and professional development.

1. IoT Architecture & Sensor Physics: Gained deep understanding of how different environmental sensors (Smoke, Flame, CO) operate, their baseline noise levels, and how to model physical thresholds in software.
2. Python Programming: Enhanced proficiency in advanced Python concepts, including Object-Oriented Programming (OOP), complex conditional logic for multi-sensor correlation, and simulation of asynchronous events.
3. Data Analytics & Visualization: Mastered the use of Matplotlib and Seaborn to transform raw environmental telemetry into professional, actionable dashboards for safety administrators.
4. Problem Solving: Developed critical thinking skills through the process of designing correlation algorithms that successfully balance the need for rapid detection against the requirement to suppress false alarms.""")

    doc.add_heading('9.2 Conclusion', level=1)
    add_content("""The Intelligent Fire Detection, Smoke Monitoring, and Emergency Warning System successfully addresses the limitations of isolated, conventional fire alarms. By leveraging IoT simulation, multi-sensor correlation (Smoke, Flame, Temp, CO), and automated response logic, the developed solution provides continuous environmental monitoring, instant alert generation, and automated mitigation (sprinklers/evacuation).

The project demonstrates that intelligent, cloud-connected fire safety systems can significantly reduce emergency response times, eliminate complacency caused by false alarms, and provide responders with critical situational awareness. The successful completion of this project fulfills all evaluation criteria and serves as a robust foundation for deploying next-generation life safety infrastructure in modern smart buildings.""")

    # Save the document
    doc.save('/home/ubuntu/Internship_Report_Fire_Detection.docx')
    print("Report generated successfully at /home/ubuntu/Internship_Report_Fire_Detection.docx")

if __name__ == "__main__":
    create_report()
